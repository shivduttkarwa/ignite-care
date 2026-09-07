import datetime as dt

import pytest
from django.urls import reverse

from apps.pdfgen.render import build_participant_book, build_record_document
from apps.records.models import AuditEvent, CareRecord, RecordStatus, Shift


def make_record(participant, author, day, shift=Shift.NIGHT):
    record = CareRecord.objects.create(
        participant=participant,
        home=participant.home,
        service_date=day,
        shift=shift,
        answers={
            "shower": True,
            "bed_bath": False,
            "physio": True,
            "breakfast": "Toast",
            "fluids": "800ml",
            "bowel": "Normal",
            "urine": "Normal",
            "sleep_from": "23:00",
            "sleep_to": "05:30",
        },
        status=RecordStatus.SUBMITTED,
        created_by=author,
        submitted_by=author,
        submitted_at=dt.datetime.combine(day, dt.time(6, 2), tzinfo=dt.UTC),
        shower=True,
        bed_bath=False,
        physio_completed=True,
        bowel_recorded=True,
        urine_recorded=True,
        fluids_recorded=True,
    )
    build_record_document(record)
    return record


def test_a_record_pdf_is_rendered_once_and_reused(db, daniel, worker):
    record = make_record(daniel, worker, dt.date(2026, 8, 12))
    first = record.document
    again = build_record_document(record)
    assert again.pk == first.pk, "a record's PDF must never be regenerated"


def test_record_book_concatenates_every_stored_pdf_oldest_first(db, daniel, worker):
    for offset in range(3):
        make_record(daniel, worker, dt.date(2026, 8, 10 + offset))

    pdf_bytes, pages = build_participant_book(daniel)
    assert pages == 3
    assert pdf_bytes.startswith(b"%PDF")


def test_book_is_empty_when_nothing_has_been_submitted(db, daniel):
    pdf_bytes, pages = build_participant_book(daniel)
    assert pages == 0
    assert pdf_bytes == b""


def test_downloading_a_record_is_audited(client, db, daniel, worker):
    record = make_record(daniel, worker, dt.date(2026, 8, 12))
    client.force_login(worker)

    response = client.get(reverse("record_pdf", args=[record.pk]))
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert AuditEvent.objects.filter(
        action=AuditEvent.Action.DOWNLOAD, target=record.reference
    ).exists()


@pytest.mark.parametrize("route", ["export_csv", "export_xlsx"])
def test_exports_are_manager_only(client, db, daniel, worker, route):
    make_record(daniel, worker, dt.date(2026, 8, 12))
    client.force_login(worker)
    response = client.get(reverse(route))
    assert response.status_code == 302, "a support worker is redirected away from exports"


def test_csv_export_carries_the_reportable_columns(client, db, daniel, worker, manager):
    make_record(daniel, worker, dt.date.today())
    client.force_login(manager)

    response = client.get(reverse("export_csv"))
    body = response.content.decode()
    header = body.splitlines()[0]

    for column in ("Physio", "Shower", "Bowel", "Fluids", "Participant", "Care worker"):
        assert column in header
    assert "Daniel Reeves" in body
    assert AuditEvent.objects.filter(action=AuditEvent.Action.EXPORT).exists()


def test_xlsx_export_opens_as_a_workbook(client, db, daniel, worker, manager):
    import io

    from openpyxl import load_workbook

    make_record(daniel, worker, dt.date.today())
    client.force_login(manager)

    response = client.get(reverse("export_xlsx"))
    book = load_workbook(io.BytesIO(response.content))
    assert book.active.title == "Care records"
    assert book.active.cell(row=1, column=1).value == "Date"


@pytest.mark.parametrize(
    ("route", "content_type", "filename"),
    [
        ("export_csv", "text/csv", "ignite-records.csv"),
        (
            "export_xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "ignite-records.xlsx",
        ),
    ],
)
def test_exports_download_rather_than_render(
    client, db, daniel, worker, manager, route, content_type, filename
):
    """Every file endpoint must be an attachment, or htmx swaps binary into the page."""
    make_record(daniel, worker, dt.date.today())
    client.force_login(manager)

    response = client.get(reverse(route))
    assert response["Content-Type"].startswith(content_type)
    assert response["Content-Disposition"] == f'attachment; filename="{filename}"'


def test_record_pdf_is_an_attachment(client, db, daniel, worker):
    record = make_record(daniel, worker, dt.date(2026, 8, 12))
    client.force_login(worker)

    response = client.get(reverse("record_pdf", args=[record.pk]))
    assert response["Content-Type"] == "application/pdf"
    assert response["Content-Disposition"].startswith("attachment;")


def test_participant_book_is_an_attachment(client, db, daniel, worker):
    make_record(daniel, worker, dt.date(2026, 8, 12))
    client.force_login(worker)

    response = client.get(reverse("participant_book", args=[daniel.pk]))
    assert response["Content-Type"] == "application/pdf"
    assert response["Content-Disposition"].startswith("attachment;")
