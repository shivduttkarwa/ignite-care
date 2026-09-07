"""Manager-facing records screen, detail drawer and exports.

Django admin is a superuser back door only. Everything a manager sees is
built here, to the approved design.
"""

import csv
import datetime as dt

from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render

from apps.pdfgen.render import build_selection_pdf
from apps.people.models import Home, Participant
from apps.records.models import AuditEvent, CareRecord, RecordStatus, Shift
from apps.records.schema import load_schema

RANGE_CHOICES = {"7": "Last 7 days", "14": "Last 14 days", "30": "Last 30 days", "all": "All time"}


def manager_required(view):
    def check(user):
        if not user.is_authenticated:
            return False
        if user.is_superuser:
            return True
        profile = getattr(user, "staff_profile", None)
        return profile is not None and profile.is_manager

    return login_required(user_passes_test(check)(view))


def _filtered(request):
    records = (
        CareRecord.objects.visible_to(request.user)
        .exclude(status=RecordStatus.DRAFT)
        .select_related("participant", "home", "submitted_by", "created_by", "document")
    )

    home = request.GET.get("home") or ""
    participant = request.GET.get("participant") or ""
    worker = request.GET.get("worker") or ""
    shift = request.GET.get("shift") or ""
    window = request.GET.get("range") or "7"

    if home:
        records = records.filter(home_id=home)
    if participant:
        records = records.filter(participant_id=participant)
    if worker:
        records = records.filter(submitted_by_id=worker)
    if shift:
        records = records.filter(shift=shift)
    if window in {"7", "14", "30"}:
        since = dt.date.today() - dt.timedelta(days=int(window) - 1)
        records = records.filter(service_date__gte=since)

    return records, {
        "home": home,
        "participant": participant,
        "worker": worker,
        "shift": shift,
        "range": window,
    }


def _active_chips(selected):
    chips = []
    if selected["home"]:
        home = Home.objects.filter(pk=selected["home"]).first()
        if home:
            chips.append({"key": "home", "label": f"Property: {home.label}"})
    if selected["participant"]:
        person = Participant.objects.filter(pk=selected["participant"]).first()
        if person:
            chips.append({"key": "participant", "label": f"Participant: {person.full_name}"})
    if selected["shift"]:
        chips.append({"key": "shift", "label": dict(Shift.choices).get(selected["shift"], "")})
    if selected["range"] != "all":
        chips.append({"key": "range", "label": RANGE_CHOICES.get(selected["range"], "")})
    return chips


@manager_required
def record_list(request):
    records, selected = _filtered(request)
    total = records.count()

    paginator = Paginator(records.order_by("-service_date", "-submitted_at"), 25)
    page = paginator.get_page(request.GET.get("page"))

    from django.contrib.auth import get_user_model

    User = get_user_model()

    return render(
        request,
        "records/record_list.html",
        {
            "page": page,
            "total": total,
            "selected": selected,
            "chips": _active_chips(selected),
            "homes": Home.objects.filter(is_active=True),
            "participants": Participant.objects.visible_to(request.user).active(),
            "workers": User.objects.filter(records_submitted__isnull=False).distinct(),
            "shifts": Shift.choices,
            "ranges": RANGE_CHOICES.items(),
            "row_ids": [row.pk for row in page.object_list],
        },
    )


@manager_required
def record_drawer(request, pk):
    record = get_object_or_404(
        CareRecord.objects.visible_to(request.user).select_related("participant", "submitted_by"),
        pk=pk,
    )
    AuditEvent.objects.create(
        actor=request.user, action=AuditEvent.Action.VIEW, target=record.reference
    )
    return render(
        request,
        "records/_drawer.html",
        {
            "record": record,
            "participant": record.participant,
            "schema": load_schema(record.schema_key, record.schema_version),
        },
    )


EXPORT_COLUMNS = [
    ("Date", lambda r: r.service_date.isoformat()),
    ("Shift", lambda r: r.get_shift_display()),
    ("Property", lambda r: r.home.name),
    ("Participant", lambda r: r.participant.full_name),
    ("Reference", lambda r: r.reference),
    ("Care worker", lambda r: (r.submitted_by or r.created_by).get_full_name()),
    ("Submitted", lambda r: r.submitted_at.isoformat() if r.submitted_at else ""),
    ("Physio", lambda r: _tri(r.physio_completed)),
    ("Shower", lambda r: _tri(r.shower)),
    ("Bed bath", lambda r: _tri(r.bed_bath)),
    ("Bowel", lambda r: "Recorded" if r.bowel_recorded else "Not recorded"),
    ("Urine", lambda r: "Recorded" if r.urine_recorded else "Not recorded"),
    ("Fluids", lambda r: (r.answers or {}).get("fluids", "")),
    ("Bowel detail", lambda r: (r.answers or {}).get("bowel", "")),
    ("Urine detail", lambda r: (r.answers or {}).get("urine", "")),
]


def _tri(value):
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return "Not recorded"


def _audit_export(request, kind, count):
    AuditEvent.objects.create(
        actor=request.user,
        action=AuditEvent.Action.EXPORT,
        target=f"records:{kind}",
        detail={"rows": count, "filters": dict(request.GET.items())},
    )


@manager_required
def export_csv(request):
    records, _selected = _filtered(request)
    records = records.order_by("service_date")
    _audit_export(request, "csv", records.count())

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="ignite-records.csv"'
    writer = csv.writer(response)
    writer.writerow([name for name, _ in EXPORT_COLUMNS])
    for record in records:
        writer.writerow([fn(record) for _, fn in EXPORT_COLUMNS])
    return response


@manager_required
def export_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    records, _selected = _filtered(request)
    records = records.order_by("service_date")
    _audit_export(request, "xlsx", records.count())

    book = Workbook()
    sheet = book.active
    sheet.title = "Care records"

    header_fill = PatternFill("solid", fgColor="115E74")
    header_font = Font(color="FFFFFF", bold=True)
    for column, (name, _fn) in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"

    for row_index, record in enumerate(records, start=2):
        for column, (_name, fn) in enumerate(EXPORT_COLUMNS, start=1):
            sheet.cell(row=row_index, column=column, value=fn(record))

    for column, (name, _fn) in enumerate(EXPORT_COLUMNS, start=1):
        letter = sheet.cell(row=1, column=column).column_letter
        sheet.column_dimensions[letter].width = max(12, min(40, len(name) + 8))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="ignite-records.xlsx"'
    book.save(response)
    return response


@manager_required
def export_pdf(request):
    ids = request.GET.getlist("id")
    records = list(
        CareRecord.objects.visible_to(request.user)
        .filter(pk__in=ids, status=RecordStatus.SUBMITTED)
        .select_related("document")
    )
    pdf_bytes, pages = build_selection_pdf(records)
    if not pages:
        return HttpResponse("None of the selected records have a stored PDF.", status=404)

    _audit_export(request, "pdf", len(records))
    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="ignite-records.pdf"'
    return response
